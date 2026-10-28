"""Loss-aware CSV/XLSX input and safe, styled spreadsheet exports."""
import csv
from pathlib import Path
import re
from decimal import Decimal, InvalidOperation

ID_HEADERS = {'订单号', '订单编号', '商户订单号', '商家订单号', '主订单编号', '主订单号', '订单id', 'order_id', 'orderid', '交易订单号', '交易号'}
AMOUNT_HEADERS = {'金额', '订单金额', '订单应收金额', '应收金额', '应收', '实收金额', '实收', '实付金额', '实付款', '付款金额', '收入金额', '收入', 'amount', 'total', '支付金额', '入账金额', '账单实收金额'}


def read_table(path):
    path = Path(path)
    if not path.is_file() or path.stat().st_size > 50 * 1024 * 1024:
        raise ValueError('请选择不超过 50MB 的表格文件')
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            rows = [list(row) for row in book.active.iter_rows(values_only=True)]
        finally:
            book.close()
    elif path.suffix.lower() in {'.csv', '.txt', '.tsv'}:
        raw = path.read_bytes()
        for encoding in ('utf-8-sig', 'gb18030', 'utf-16'):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeError:
                continue
        else:
            raise ValueError('文件编码无法识别，请另存为 UTF-8 CSV')
        if not text.strip():
            raise ValueError('表格为空')
        import io
        delimiter = '\t' if path.suffix.lower() == '.tsv' or '\t' in text.splitlines()[0] else ','
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    else:
        raise ValueError('支持 CSV / TSV / XLSX；旧 XLS 请先另存为 XLSX')
    if not rows:
        raise ValueError('表格为空')
    header_at = 0
    for i, row in enumerate(rows[:100]):
        names = {str(c).strip().lower() for c in row if c is not None}
        if names & ID_HEADERS and names & AMOUNT_HEADERS:
            header_at = i
            break
    headers = [str(v).strip() if v is not None else '' for v in rows[header_at]]
    if len(headers) > 200 or len(set(headers)) != len(headers):
        raise ValueError('表头不得重复，列数不得超过 200')
    result = []
    for number, row in enumerate(rows[header_at+1:], header_at+2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        record = {key: (row[i].strip().replace('\r\n', '\n') if isinstance(row[i], str) else row[i]) if i < len(row) else ''
                  for i, key in enumerate(headers) if key}
        record['__row__'] = number
        result.append(record)
        if len(result) > 100000:
            raise ValueError('当前版本单表最多支持 100000 行')
    return result


def normalize_order_id(value):
    if value is None or isinstance(value, bool):
        raise ValueError('订单号为空或格式无效')
    numeric_cell = isinstance(value, (int, float))
    text = str(value).strip().lstrip('\ufeff')
    if text.startswith("'"):
        text = text[1:]
    if not text or text.startswith(('=', '+', '@')):
        raise ValueError('订单号不能是公式或空值')
    if re.fullmatch(r'\d+(?:\.\d+)?(?:[eE][+-]?\d+)?', text):
        try:
            number = Decimal(text)
            if number.adjusted() > 63 or number != number.to_integral_value():
                raise ValueError('订单号必须为整数，且不超过 64 位')
            if 'e' in text.lower() or '.' in text:
                normalized = format(number, '.0f')
            else:
                normalized = text
            issue = ''
            if numeric_cell and len(normalized) > 15:
                issue = '长订单号曾以数字存储，原始尾数需核实'
            elif 'e' in text.lower() and len(normalized) > 15 and len(number.as_tuple().digits) < len(normalized):
                issue = '科学计数法有效位数不足，原始尾数需核实'
            return normalized, issue
        except InvalidOperation as error:
            raise ValueError('订单号格式错误') from error
    if len(text) > 64 or re.search(r'[\x00-\x1f]', text):
        raise ValueError('订单号过长或含控制字符')
    return text, ''


def find_column(rows, aliases, requested=''):
    if not rows:
        return requested
    if requested and requested != '自动':
        if requested not in rows[0]:
            raise ValueError('找不到指定列：' + requested)
        return requested
    for key in rows[0]:
        if key.strip().lower() in aliases:
            return key
    raise ValueError('无法识别列，请在列映射处填写实际表头：' + ' / '.join(sorted(aliases)[:5]))


def display_value(value):
    if value is None:
        return ''
    return str(value)


def safe_csv(value):
    text = display_value(value)
    # Spreadsheet formulas and long numeric identifiers must remain literal text.
    stripped = text.lstrip()
    if stripped.startswith(('=', '+', '-', '@', '\t', '\r')) or text.startswith(('\t', '\r', '\n')):
        return "'" + text
    if re.fullmatch(r'\d{16,}', text) or re.fullmatch(r'0\d+', text):
        return "'" + text
    return text


def write_rows(path, rows, guard=None):
    rows = guard.export_rows(rows) if guard is not None else list(rows)
    if not rows:
        raise ValueError('没有可导出的结果')
    path = Path(path)
    headers = list(dict.fromkeys(k for row in rows for k in row if not k.startswith('__')))
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == '.csv':
        with path.open('w', encoding='utf-8-sig', newline='') as handle:
            writer = csv.writer(handle)
            writer.writerow([safe_csv(k) for k in headers])
            writer.writerows([[safe_csv(row.get(k, '')) for k in headers] for row in rows])
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        book = Workbook()
        sheet = book.active
        sheet.title = '核对与分析结果'
        for row_i, values in enumerate([dict(zip(headers, headers))] + rows, 1):
            for col_i, key in enumerate(headers, 1):
                cell = sheet.cell(row_i, col_i, display_value(values.get(key, '')))
                cell.data_type = 's'
                cell.number_format = '@'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if row_i == 1:
                    cell.fill = PatternFill('solid', fgColor='123B46')
                    cell.font = Font(color='FFFFFF', bold=True)
                elif values.get('异常') == '是' or values.get('级别') in {'高危', '严重'}:
                    cell.fill = PatternFill('solid', fgColor='FDE2E2')
            if row_i > 1:
                sheet.row_dimensions[row_i].height = 34
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(48, max(16, len(str(col[0].value))*2+4))
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        book.save(path)
        book.close()
    else:
        raise ValueError('请选择 .csv 或 .xlsx 导出格式')
    return len(rows)
