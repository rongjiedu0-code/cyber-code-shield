import csv
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import pytest
from toolbox.common import money, integer
from toolbox.tables import normalize_order_id, read_table, write_rows, ID_HEADERS, AMOUNT_HEADERS, find_column, safe_csv
from toolbox.reconcile import reconcile_rows
from toolbox.license_guard import LicenseGuard, validate_key, machine_code, clean_machine
from key_fixtures import generate
from toolbox.radar import pain_score, analyze_signals, score_four_dimensions, collect_community, PAIN_WORDS
from toolbox.trends import parse_trends, collect_trends, analyze_products
from toolbox.security import scan_code, scan_path, redact_all
from toolbox.urls import clean_url
from toolbox.service import ToolboxService


@pytest.mark.parametrize('raw,want', [('1','1.00'),('1.005','1.01'),('0.105','0.11'),('-1.005','-1.01'),
    ('1,000.20','1000.20'),('¥9.99','9.99'),('￥9.99','9.99'),('CNY12','12.00'),
    ('RMB12','12.00'),('(9.20)','-9.20'),('0','0.00'),('0.00','0.00'),(2,'2.00'),
    (Decimal('1.1'),'1.10'),('  2.2  ','2.20'),('1e2','100.00')])
def test_decimal_money(raw,want):
    assert str(money(raw))==want


@pytest.mark.parametrize('raw',[None,True,False,'',' ', 'NaN','Infinity','-Infinity','12,3','1,23,456','abc','1e90'])
def test_invalid_money(raw):
    with pytest.raises(ValueError):money(raw)


@pytest.mark.parametrize('value',[-1,-.01,'-12','(23)'])
def test_nonnegative_prices(value):
    with pytest.raises(ValueError):money(value,allow_negative=False)


@pytest.mark.parametrize('raw,want', [('000123','000123'),("'000123",'000123'),(' 123 ','123'),
    ('123456789012345678','123456789012345678'),('1.23456789012345678E+17','123456789012345678'),
    ('12.0','12'),('OD-2026-0001','OD-2026-0001'),('订单甲01','订单甲01'),(123,'123'),
    ('\ufeff0099','0099'),('9.001E+3','9001')])
def test_exact_identifiers(raw,want):
    normalized,issue=normalize_order_id(raw)
    assert normalized==want and not issue


@pytest.mark.parametrize('raw',['1.23E+17',1.2345678901234568e17,123456789012345678,'1.23456789012345E17'])
def test_lost_digits_are_flagged(raw):
    assert normalize_order_id(raw)[1]


@pytest.mark.parametrize('raw',[None,True,False,'','  ','=123','+123','@foo','1.5','1e100','x'*65,'a\nb'])
def test_invalid_identifiers(raw):
    with pytest.raises(ValueError):normalize_order_id(raw)


@pytest.mark.parametrize('alias',sorted(ID_HEADERS))
def test_order_column_aliases(alias):
    assert find_column([{alias:'1'}],ID_HEADERS)==alias


@pytest.mark.parametrize('alias',sorted(AMOUNT_HEADERS))
def test_amount_column_aliases(alias):
    assert find_column([{alias:'1'}],AMOUNT_HEADERS)==alias


@pytest.mark.parametrize('value',[True,1.0,'1',0,-1,101,None])
def test_limit_validation(value):
    with pytest.raises(ValueError):integer(value,'limit',1,100)


@pytest.mark.parametrize('machine',['DRJ88A01','12345678','ABCDEFGH','00000000']+[f'{i:08X}' for i in range(20)])
@pytest.mark.parametrize('mode',['web','sha'])
def test_both_license_tracks(machine,mode):
    key=generate(machine,mode)
    assert validate_key(machine,key)
    assert not validate_key('ZZZZZZZZ' if machine!='ZZZZZZZZ' else 'YYYYYYYY',key)


@pytest.mark.parametrize('bad',['','a','1234567','123456789','ABCD!FGH','中文机器码123','AB CD123','abc-1234'])
def test_machine_rejects_invalid(bad):
    with pytest.raises(ValueError):clean_machine(bad)


@pytest.mark.parametrize('code',['','VIP-0000-0000-0000','BAD-1234-5678-9999','VIP-123-1234-1234',None,'VIP-GGGG-0000-0000'])
def test_bad_keys_fail(code):
    assert not validate_key('DRJ88A01',code)


def test_persisted_activation_and_wrong_machine(tmp_path):
    p=tmp_path/'license.json'
    a=LicenseGuard('DRJ88A01',p)
    assert not a.activated
    a.activate(generate(a.machine,'sha'))
    assert LicenseGuard(a.machine,p).activated
    assert not LicenseGuard('DRJ88A02',p).activated
    p.write_text('{invalid',encoding='utf-8')
    assert not a.activated


def test_machine_identity_is_stable_and_uppercase():
    assert machine_code()==machine_code()
    assert len(machine_code())==8 and machine_code().upper()==machine_code()


@pytest.mark.parametrize('count',[1,5,10,11,30,51])
@pytest.mark.parametrize('extension',['.csv','.xlsx'])
@pytest.mark.parametrize('active',[False,True])
def test_trial_gate_every_export(tmp_path,count,extension,active):
    guard=LicenseGuard('DRJ88A01',tmp_path/'license.json')
    if active:guard.activate(generate(guard.machine))
    target=tmp_path/('result'+extension)
    rows=[{'订单号':str(123456789012345678+i),'异常':'是' if i==1 else '否','内容':'=1+1'} for i in range(count)]
    written=write_rows(target,rows,guard)
    assert written==(count if active else min(10,count))
    if extension=='.xlsx':
        from openpyxl import load_workbook
        b=load_workbook(target);s=b.active
        assert s.max_row==written+1
        assert s['A2'].value==rows[0]['订单号'] and s['A2'].data_type=='s'
        assert s['C2'].data_type=='s'
        if written>1:assert s['A3'].fill.fgColor.rgb.endswith('FDE2E2')
        b.close()
    else:
        with target.open(encoding='utf-8-sig',newline='') as f:data=list(csv.DictReader(f))
        assert len(data)==written and data[0]['订单号'].startswith("'") and data[0]['内容'].startswith("'")


@pytest.mark.parametrize('payload',['=1+1',' +SUM(A1:A2)','-1+2','@cmd','\t=2','\r=3','000001','123456789012345678'])
def test_csv_literal_protection(payload):
    assert safe_csv(payload).startswith("'")


def test_reconcile_all_anomalies_and_exact_decimals():
    order=[{'订单号':'001','金额':'10.10'},{'订单号':'002','金额':'20'}, {'订单号':'003','金额':'5'}, {'订单号':'005','金额':'1'}, {'订单号':'005','金额':'2'}]
    payment=[{'订单号':'001','金额':'10.10'},{'订单号':'002','金额':'8'},{'订单号':'002','金额':'7'}, {'订单号':'004','金额':'9'},{'订单号':'005','金额':'3'}]
    rows={r['订单号']:r for r in reconcile_rows(order,payment)}
    assert rows['001']['状态']=='核对一致'
    assert '多笔入账' in rows['002']['状态'] and '金额不一致' in rows['002']['状态']
    assert rows['002']['差额']=='-5.00'
    assert rows['003']['状态']=='漏收款' and rows['004']['状态']=='无源流水'
    assert '重复订单' in rows['005']['状态']


def test_reconcile_refund_negative_and_precision_rejection():
    rows=reconcile_rows([{'订单号':'1','金额':'-10'}],[{'订单号':'1','金额':'-10'}])
    assert rows[0]['状态']=='核对一致'
    rows=reconcile_rows([{'订单号':'1.23E17','金额':'1'}],[{'订单号':'123000000000000000','金额':'1'}])
    assert any(r['状态']=='输入待核实' for r in rows)
    assert not any(r['状态']=='核对一致' for r in rows)


def test_manual_mapping_and_invalid_amount():
    rows=reconcile_rows([{'单':'A','价':'bad'}],[{'号':'A','钱':'2'}],{'order_id':'单','order_amount':'价','payment_id':'号','payment_amount':'钱'})
    assert any(r['状态']=='输入待核实' for r in rows)
    with pytest.raises(ValueError):reconcile_rows([{'abc':'1'}],[])


@pytest.mark.parametrize('encoding',['utf-8-sig','gb18030','utf-16'])
def test_csv_encodings_and_preamble(tmp_path,encoding):
    p=tmp_path/'a.csv';p.write_text('账单说明\n订单号,金额\n0001,2.20\n',encoding=encoding)
    assert read_table(p)[0]['订单号']=='0001'


def test_csv_multiline_and_duplicate_header(tmp_path):
    p=tmp_path/'a.csv';p.write_text('标题,内容\n"第一条","两行\n反馈"\n',encoding='utf-8')
    assert read_table(p)[0]['内容']=='两行\n反馈'
    p.write_text('订单号,订单号\n1,2',encoding='utf-8')
    with pytest.raises(ValueError):read_table(p)
    p.write_text('',encoding='utf-8')
    with pytest.raises(ValueError):read_table(p)


@pytest.mark.parametrize('phrase',list(PAIN_WORDS))
def test_pain_intent_words(phrase):
    score,matches=pain_score('客户说：'+phrase)
    assert score>20 and phrase in matches


@pytest.mark.parametrize('values,want', [((10,10,10,10),100),((0,0,0,0),0),((8,8,8,8),80),((5,5,5,5),50),((9,8.5,9,7.5),85.75)])
def test_original_four_dimension_weights(values,want):
    assert score_four_dimensions(*values)[0]==round(want,1)


def test_radar_deduplicates_links_and_labels_evidence():
    records=[{'title':'手动对账很麻烦','url':'https://example.cn/post?id=1&utm_source=a'},
             {'title':'不同标题','url':'https://example.cn/post?id=1'}, {'title':'普通标题'}]
    rows=analyze_signals(records)
    assert len(rows)==2 and rows[0]['商机分']>rows[1]['商机分']
    assert '尚无需求证据' in rows[1]['判读']
    assert len(analyze_signals(records,'对账'))==1


def test_community_real_parser_boundary():
    rows=collect_community(get_json=lambda _:dict(code=0,data={'list':[dict(title='每天手动对账',desc='耗时',bvid='BV1234567890')]}))
    assert len(rows)==1 and rows[0]['来源链接'].startswith('https://www.bilibili.com/video/')
    with pytest.raises(ValueError):collect_community(get_json=lambda _:dict(code=-1))


@pytest.mark.parametrize('source,payload', [('百度',{'success':True,'data':{'cards':[{'content':[{'word':'收纳盒','hotScore':'123'}]}]}}),
    ('哔哩哔哩',{'code':0,'list':[{'keyword':'收纳盒'}]}),('今日头条',{'data':[{'Title':'收纳盒','Url':'https://www.toutiao.com/a1?utm_source=a'}]})])
def test_live_trend_parsers(source,payload):
    rows=parse_trends(source,payload)
    assert rows[0]['热词']=='收纳盒' and rows[0]['排名']==1
    assert '价格' not in rows[0] and 'utm_' not in rows[0]['来源链接']


@pytest.mark.parametrize('source,payload',[('百度',{'success':False}),('哔哩哔哩',{'code':-1}),('今日头条',{'data':None}),('未知',{})])
def test_invalid_trend_payload(source,payload):
    with pytest.raises(ValueError):parse_trends(source,payload)


def test_empty_success_is_not_fake_fallback():
    assert parse_trends('哔哩哔哩',{'code':0,'list':[]})==[]
    assert collect_trends(get_json=lambda _:{'success':True,'data':{'cards':[]}})==[]
    def fail(_):raise OSError('网络中断')
    with pytest.raises(RuntimeError):collect_trends(get_json=fail)


@pytest.mark.parametrize('rate',['0.01','1%'])
def test_real_product_profit_and_price_band(rate):
    rows=analyze_products([{'商品名称':'a','销售价':'100','采购成本':'60','运费':'5','手续费率':rate,'品类':'收纳'},
                           {'商品名称':'b','销售价':'120','采购成本':'90','品类':'收纳'}])
    assert rows[0]['估算毛利']=='34.00' and rows[0]['估算毛利率']=='34.00%'
    assert rows[0]['同表品类价格带']=='100.00–120.00' and rows[0]['同表品类中位价']=='110.00'


@pytest.mark.parametrize('sale,rate,expected_fee',[('19.50','3%','0.59'),('1.00','0.5%','0.01')])
def test_half_cent_fees_round_consistently(sale,rate,expected_fee):
    row=analyze_products([{'商品名称':'边界核对','销售价':sale,'采购成本':'0','手续费率':rate}])[0]
    assert row['手续费']==expected_fee
    assert Decimal(row['估算毛利'])==Decimal(sale)-Decimal(expected_fee)


@pytest.mark.parametrize('bad',[-1,'NaN','Infinity','2','-1%'])
def test_invalid_fee_rates(bad):
    with pytest.raises((ValueError,ArithmeticError)):analyze_products([{'商品名称':'a','销售价':'1','采购成本':'1','手续费率':bad}])


@pytest.mark.parametrize('content,rule',[
    ('import subprocess\nsubprocess.run(x,shell=True)','CCS-AST-009'),
    ('import os\nos.system(user)','CCS-EXEC-006'),('eval(user)','CCS-EXEC-006'),
    ('import pickle\npickle.loads(data)','CCS-DESERIALIZE-007'),
    ('password = "private-example-12345"','CN-SECRET-011'),
    ('child_process.exec(userInput);','CN-JS-012'),('db.execute(f"SELECT {user}")','CN-SQL-013'),
    ('(base / name).read_text()','CCS-AST-010'),('const x = new Function(input);','CN-JS-012')])
def test_security_patterns(content,rule):
    rows=scan_code(content)
    assert any(r['规则']==rule for r in rows)
    assert 'private-example-12345' not in json.dumps(rows)


def test_security_never_executes_and_skips_dependencies(tmp_path):
    sentinel=tmp_path/'MUST_NOT_EXIST'
    (tmp_path/'a.py').write_text(f'from pathlib import Path\nPath({str(sentinel)!r}).write_text("executed")\neval(user)',encoding='utf-8')
    excluded=tmp_path/'node_modules';excluded.mkdir();(excluded/'a.js').write_text('eval(x)')
    rows=scan_path(tmp_path)
    assert not sentinel.exists() and all('node_modules' not in r['文件'] for r in rows)


@pytest.mark.parametrize('text',['api_key="private-example-123"','secret = "private-example-123"','password: "private-example-123"'])
def test_report_redacts_assignments(text):
    assert 'private-example-123' not in redact_all(text)


@pytest.mark.parametrize('parameter',['utm_source','utm_medium','utm_campaign','fbclid','gclid','spm','scm','share_token','xsec_source','ref'])
def test_url_tracking_removed(parameter):
    assert clean_url(f'https://example.cn/p?sku=2&{parameter}=tracking#frag')=='https://example.cn/p?sku=2'


@pytest.mark.parametrize('url',['file:///a','ftp://x/a','https://user:pass@example.cn','https://example.cn/ bad','not-a-url'])
def test_invalid_urls(url):
    with pytest.raises(ValueError):clean_url(url)


def test_service_export_cannot_bypass_gate(tmp_path):
    guard=LicenseGuard('DRJ88A01',tmp_path/'activation.json')
    service=ToolboxService(guard)
    service.rows=[{'结果':i} for i in range(35)]
    assert len(service.state()['preview'])==10
    assert service.export(tmp_path/'trial.csv')==10
    guard.activate(generate(guard.machine))
    assert service.export(tmp_path/'full.xlsx')==35
    assert len(service.state()['preview'])==35
